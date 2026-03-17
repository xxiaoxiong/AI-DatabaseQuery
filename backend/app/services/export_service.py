import io
import csv
from datetime import datetime
from typing import Optional
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


async def export_to_excel(
    columns: list[str],
    rows: list[dict],
    datasource_name: str,
    question: str,
    sql: str,
) -> bytes:
    """导出查询结果为 Excel 文件"""
    wb = Workbook()
    ws = wb.active
    ws.title = "查询结果"

    # 添加元数据行
    ws['A1'] = "数据源"
    ws['B1'] = datasource_name
    ws['A2'] = "查询时间"
    ws['B2'] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    ws['A3'] = "问题"
    ws['B3'] = question
    ws['A4'] = "SQL"
    ws['B4'] = sql

    # 设置元数据行样式
    header_fill = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")
    header_font = Font(bold=True)
    for row in range(1, 5):
        ws[f'A{row}'].fill = header_fill
        ws[f'A{row}'].font = header_font

    # 添加数据表头
    header_row = 6
    for col_idx, col_name in enumerate(columns, 1):
        cell = ws.cell(row=header_row, column=col_idx)
        cell.value = col_name
        cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # 添加数据行
    for row_idx, row in enumerate(rows, header_row + 1):
        for col_idx, col_name in enumerate(columns, 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            value = row.get(col_name)
            cell.value = value
            cell.alignment = Alignment(horizontal="left", vertical="center")

    # 自动调整列宽
    for col_idx, col_name in enumerate(columns, 1):
        max_length = len(str(col_name))
        for row in rows:
            cell_value = str(row.get(col_name, ""))
            max_length = max(max_length, len(cell_value))
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[get_column_letter(col_idx)].width = adjusted_width

    # 保存到字节流
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()


async def export_to_csv(
    columns: list[str],
    rows: list[dict],
    datasource_name: str,
    question: str,
    sql: str,
) -> bytes:
    """导出查询结果为 CSV 文件"""
    output = io.StringIO()
    writer = csv.writer(output)

    # 添加元数据行
    writer.writerow(["数据源", datasource_name])
    writer.writerow(["查询时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")])
    writer.writerow(["问题", question])
    writer.writerow(["SQL", sql])
    writer.writerow([])  # 空行分隔

    # 添加表头
    writer.writerow(columns)

    # 添加数据行
    for row in rows:
        writer.writerow([row.get(col, "") for col in columns])

    return output.getvalue().encode('utf-8-sig')  # UTF-8 with BOM for Excel

